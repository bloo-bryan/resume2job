"""Generate PDF and DOCX fixture files for parser integration tests."""

from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "data" / "samples" / "fixtures"

RESUME_TEXT = """\
Sarah Chen
Email: sarah.chen@email.com | Phone: (555) 123-4567 | LinkedIn: linkedin.com/in/sarahchen

Summary
Senior Machine Learning Engineer with 8+ years of experience building production ML systems.
Skilled in Python, PyTorch, TensorFlow, and cloud deployment.

Experience

Senior ML Engineer
Acme AI Corp
Jan 2020 - Present
- Led development of NLP pipeline processing 1M+ documents daily
- Built recommendation engine using deep learning, improving CTR by 35%
- Deployed models on AWS using Docker and Kubernetes

ML Engineer
DataTech Solutions
Jun 2016 - Dec 2019
- Developed computer vision models for defect detection using PyTorch
- Built data pipelines with Python, SQL, and Apache Spark
- Implemented A/B testing framework for model evaluation

Education
M.S. in Computer Science from Stanford University, 2016
B.S. in Mathematics from UC Berkeley, 2014

Skills
Python, PyTorch, TensorFlow, scikit-learn, SQL, Docker, Kubernetes, AWS, Git,
machine learning, deep learning, NLP, computer vision, data analysis
"""


def generate_pdf() -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    for line in RESUME_TEXT.strip().split("\n"):
        pdf.cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")

    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    pdf.output(str(FIXTURES_DIR / "sample_resume.pdf"))
    print(f"Generated: {FIXTURES_DIR / 'sample_resume.pdf'}")


def generate_docx() -> None:
    import docx

    doc = docx.Document()
    for line in RESUME_TEXT.strip().split("\n"):
        doc.add_paragraph(line)

    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(FIXTURES_DIR / "sample_resume.docx"))
    print(f"Generated: {FIXTURES_DIR / 'sample_resume.docx'}")


if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    print("Done.")
